from datetime import datetime
import email.mime.multipart
import email.mime.text
import base64
import os
import smtplib
import openpyxl
import requests


def get_job_content_automatic():
  schedule_file_path = "Lich lam viec 2026.xlsx"
  task_file_path = "Cong viec ca dem.xlsx"

  today = datetime.now()
  current_day = today.day
  current_month = today.month

  # Kiểm tra xem file Excel lịch làm việc có tồn tại không
  if not os.path.exists(schedule_file_path):
    raise FileExistsError(
        f"Không tìm thấy file lịch làm việc '{schedule_file_path}'. Vui lòng"
        " upload file lên GitHub!"
    )

  # 1. Đọc file lịch làm việc
  wb = openpyxl.load_workbook(schedule_file_path, data_only=True)
  target_sheet = next(
      (
          wb[name]
          for name in wb.sheetnames
          if str(current_month) in name or f"0{current_month}" in name
      ),
      wb.active,
  )

  header_row_idx = 7
  target_column_index = None

  for col_idx in range(1, target_sheet.max_column + 1):
    cell_val = target_sheet.cell(row=header_row_idx, column=col_idx).value
    if cell_val is not None:
      if isinstance(cell_val, datetime):
        if cell_val.day == current_day and cell_val.month == current_month:
          target_column_index = col_idx
          break
      else:
        val_str = str(cell_val).strip()
        if (
            val_str in (str(current_day), f"0{current_day}")
            or f"{current_day}/{current_month}" in val_str
        ):
          target_column_index = col_idx
          break

  if target_column_index is None:
    raise ValueError(
        f"Không tìm thấy cột ngày {current_day}/{current_month} trong file lịch"
        " làm việc!"
    )

  # Tìm tên Trưởng ca (Cột 10 & 11)
  name_col_idx, shift_leader_col, start_row = 10, 11, 8
  duty_persons = []

  for r in range(start_row, target_sheet.max_row + 1):
    cell_val = (
        str(target_sheet.cell(row=r, column=target_column_index).value)
        .strip()
        .upper()
    )
    col_k_val = (
        str(target_sheet.cell(row=r, column=shift_leader_col).value)
        .strip()
        .lower()
    )

    if cell_val == "N1" and "shift leader" in col_k_val:
      person_name = str(
          target_sheet.cell(row=r, column=name_col_idx).value
      ).strip()
      if person_name and person_name != "None":
        duty_persons.append(person_name)

  leader_str = ", ".join(duty_persons) if duty_persons else "Không tìm thấy"

  # 2. Đọc file công việc ca đêm (Cột B là ngày, Cột C là nội dung, bắt đầu từ dòng 3)
  night_tasks = []
  if os.path.exists(task_file_path):
    try:
      wb_task = openpyxl.load_workbook(task_file_path, data_only=True)
      sheet_task = wb_task.active
      for r in range(3, sheet_task.max_row + 1):
        cell_day_val = sheet_task.cell(row=r, column=2).value
        cell_task_val = sheet_task.cell(row=r, column=3).value
        if cell_day_val is not None:
          try:
            if int(float(cell_day_val)) == current_day:
              if cell_task_val and str(cell_task_val).strip() != "None":
                night_tasks.append(str(cell_task_val).strip())
          except:
            pass
    except Exception as e:
      print(f"Lỗi đọc file task: {e}")

  if not night_tasks:
    night_tasks = ["Vệ sinh trạm điện định kỳ theo kế hoạch"]

  # Form chuẩn yêu cầu
  result_lines = [
      "Đội điện",
      f"Trưởng ca: {leader_str}",
  ]
  result_lines.extend([f"- {task}" for task in night_tasks])

  return "\n".join(result_lines)


def save_history_to_github(message):
  """Tự động lưu lịch sử báo cáo thành file .txt vào thư mục history/ trên GitHub"""
  token = os.getenv("GITHUB_TOKEN")
  repo = os.getenv("GITHUB_REPOSITORY")  # dạng owner/repo
  if not token or not repo:
    print("Thiếu GITHUB_TOKEN hoặc GITHUB_REPOSITORY, bỏ qua lưu lịch sử.")
    return

  today_str = datetime.now().strftime("%Y-%m-%d")
  file_path = f"history/{today_str}.txt"
  url = f"https://api.github.com/repos/{repo}/contents/{file_path}"

  headers = {
      "Authorization": f"Bearer {token}",
      "Accept": "application/vnd.github+json",
  }

  # Kiểm tra xem file ngày hôm nay đã tồn tại trên kho chưa
  response = requests.get(url, headers=headers)
  sha = None
  if response.status_code == 200:
    sha = response.json().get("sha")

  # Mã hóa nội dung sang Base64 để đẩy lên GitHub API
  content_bytes = message.encode("utf-8")
  content_base64 = base64.b64encode(content_bytes).decode("utf-8")

  data = {
      "message": f"Auto-save report for {today_str}",
      "content": content_base64,
      "branch": "main",  # hoặc master tùy nhánh của bạn
  }
  if sha:
    data["sha"] = sha

  put_resp = requests.put(url, headers=headers, json=data)
  if put_resp.status_code in [200, 201]:
    print(f"Đã tự động lưu lịch sử báo cáo vào {file_path} thành công!")
  else:
    print(
        f"Lỗi khi lưu lịch sử lên GitHub: {put_resp.status_code} -"
        f" {put_resp.text}"
    )


def send_telegram(message):
  token = os.getenv("TELEGRAM_TOKEN")
  chat_id = os.getenv("TELEGRAM_CHAT_ID")
  if not token or not chat_id:
    return

  url = f"https://api.telegram.org/bot{token}/sendMessage"
  md_message = message.replace("Đội điện", "*Đội điện*").replace(
      "Trưởng ca:", "*Trưởng ca:*"
  )

  payload = {"chat_id": chat_id, "text": md_message, "parse_mode": "Markdown"}
  requests.post(url, json=payload)


def send_email(message, subject_title="Báo cáo ca trực"):
  sender_email = os.getenv("GMAIL_USER")
  app_password = os.getenv("GMAIL_APP_PASS")

  if not sender_email or not app_password:
    return

  subject = (
      f"{subject_title} - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
  )
  msg = email.mime.multipart.MIMEMultipart()
  msg["From"] = sender_email
  msg["To"] = sender_email
  msg["Subject"] = subject
  msg.attach(email.mime.text.MIMEText(message, "plain", "utf-8"))

  try:
    server = smtplib.SMTP("smtp.gmail.com", 587)
    server.starttls()
    server.login(sender_email, app_password)
    server.sendmail(sender_email, sender_email, msg.as_string())
    server.quit()
  except Exception as e:
    print(f"Lỗi gửi Email: {e}")


if __name__ == "__main__":
  try:
    # Lấy nội dung báo cáo tự động
    msg = get_job_content_automatic()

    # Gửi qua Telegram và Gmail
    send_telegram(msg)
    send_email(msg, subject_title="Báo cáo ca trực")

    # Tự động lưu lịch sử vào kho GitHub
    save_history_to_github(msg)

  except Exception as e:
    # TÍNH NĂNG CẢNH BÁO LỖI: Nếu có sự cố file Excel, gửi thông báo đỏ ngay lập tức
    error_msg = (
        f"🚨 CẢNH BẢO LỖI BOT BÁO CÁO 🚨\nThời gian: {datetime.now()}\nChi tiết"
        f" lỗi: {str(e)}"
    )
    print(error_msg)
    send_telegram(error_msg)
    send_email(error_msg, subject_title="[LỖI] Báo cáo ca trực thất bại")
